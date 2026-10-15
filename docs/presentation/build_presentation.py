"""Build the six-slide HTML, PDF and editable deck. No app or provider calls."""
from pathlib import Path
import html
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as Shape
from pptx.util import Inches, Pt
from playwright.sync_api import sync_playwright
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
BLUE, NAVY, INK, SOFT, BORDER = "0052FF", "001F61", "000B24", "F0F5FF", "C7D1D9"
ICONS = {
    "risk": '<path d="M32 8 57 54H7Z"/><path d="M32 24v13m0 7v2"/>',
    "evidence": '<path d="M14 7h29l8 9v41H14Z M23 23h19m-19 10h19m-19 10h12"/>',
    "person": '<circle cx="32" cy="20" r="11"/><path d="M12 57v-8a20 20 0 0 1 40 0v8"/>',
    "browser": '<rect x="5" y="10" width="54" height="43" rx="5"/><path d="M5 23h54M13 17h2m6 0h2"/>',
    "shield": '<path d="m32 6 23 10v18c-3 12-12 20-23 25C21 54 12 46 9 34V16Z M21 32l8 8 15-18"/>',
    "database": '<ellipse cx="32" cy="13" rx="23" ry="8"/><path d="M9 13v38c0 11 46 11 46 0V13M9 31c0 11 46 11 46 0"/>',
    "spark": '<path d="m32 5 7 20 20 7-20 7-7 20-7-20L5 32l20-7Z"/>',
    "check": '<circle cx="32" cy="32" r="25"/><path d="m17 32 11 11 20-23"/>',
    "copy": '<rect x="20" y="19" width="35" height="39" rx="4"/><path d="M42 19V7H8v39h12"/>',
    "chart": '<path d="M8 7v49h49M18 47V33h7v14m9 0V21h7v26m9 0V10h7v37"/>',
}


def icon(name, small=False):
    return f'<svg class="icon{" small" if small else ""}" viewBox="0 0 64 64" aria-hidden="true">{ICONS[name]}</svg>'


def card(name, title, text=""):
    return f'<article class="card">{icon(name)}<h3>{title}</h3><p>{text}</p></article>'


def step(name, title):
    return f'<div class="step">{icon(name,True)}<strong>{title}</strong></div>'


DATA = [
    ("THE DECISION", "Know what matters.<br>Before the conversation.",
     "Turn scattered records into a clear renewal discussion. Review evidence, resolve conflicts, and keep the producer in control.",
     '<div class="cards three">'+card("risk","See the risk")+card("evidence","Check the evidence")+card("person","Choose the next step")+'</div>'),
    ("AUTHORIZED DATA", "Real records. Authorized scope.", "",
     '<div class="metrics"><article><b>9</b><span>Authorized clients</span></article><article><b>2</b><span>Producers</span></article><article><b>$5.44M</b><span>Annual premium</span></article></div>'
     '<div class="producer-row"><div><strong>Jordan Davis</strong><span>5 clients · $3.33M</span></div><div><strong>Priya Nair</strong><span>4 clients · $2.11M</span></div></div>'
     '<p class="caption">10 source clients; Harbor excluded by access rules.</p><p class="disclosure">Synthetic assumptions: 12% commission; 35% paid loss; 100% retained demo outcomes.</p>'),
    ("ARCHITECTURE", "Simple boundaries. Clear responsibilities.",
     "Permissions and typed contracts protect every boundary. Business records stay relational; AI produces a reviewable brief.",
     '<div class="architecture"><div class="lane">'+step("browser","React + TypeScript")+'<span class="arrow">→</span>'+step("shield","FastAPI + Pydantic")+'<span class="arrow">→</span>'+step("database","SQLite")+'</div><div class="branch"><span>↓</span>'+step("spark","Configured models")+'<span>↑</span></div></div>'),
    ("GROUNDED AI", "Evidence travels with the answer.",
     "Retrieve permitted evidence, validate structured output, and resolve canonical citations. If AI is unavailable, show a labeled source-based fallback.",
     '<div class="flow">'+'<span class="arrow">→</span>'.join(step(name,label) for name,label in [("shield","Permissions"),("database","SQL FTS5"),("evidence","Typed context"),("spark","Models"),("check","Citations"),("copy","Brief")])+'</div><p class="caption">User and source hashes prevent stale reuse.</p>'),
    ("CURRENT EXPERIENCE", "Prepare clearly. Act deliberately.", "",
     '<div class="cards two">'+card("person","Producer","Select client. Read the brief. Inspect sources. Edit and copy a draft manually.")+card("chart","Regional Leader","Prioritize renewals. Open client preparation. Review the database-backed dashboard.")+'</div><p class="disclosure">No autonomous messages, policy changes, or coverage binding.</p>'),
    ("EVIDENCE & DEMO", "A working foundation. Honest limits.", "",
     '<div class="metrics"><article><b>2.37s</b><span>Observed live response</span></article><article><b>1,460</b><span>Measured tokens</span></article><article><b>C05 + E05</b><span>Canonical citations</span></article></div>'
     '<p class="caption">One Pinecrest request; no latency guarantee. Synthetic data, two personas, no proven business uplift.</p>'
     '<div class="demo">Demo: prioritize <span>→</span> inspect <span>→</span> prepare <span>→</span> copy</div>'),
]

CSS = """
:root{--blue:#0052ff;--navy:#001f61;--ink:#000b24;--body:#29343d;--soft:#f0f5ff;--border:#c7d1d9}*{box-sizing:border-box}
body{margin:0;font:16px Arial,sans-serif;color:var(--ink);background:#f0f5ff}.toolbar,.status{display:flex;align-items:center;justify-content:space-between;gap:12px;padding:14px 24px;background:white}.toolbar{border-bottom:1px solid var(--border)}.brand{color:var(--navy);font-size:13px;font-weight:700;letter-spacing:.1em}.controls{display:flex;gap:8px}button,a{min-height:40px;padding:10px 14px;border:1px solid var(--border);border-radius:8px;color:var(--navy);background:white;font:inherit;text-decoration:none;cursor:pointer}button.primary{background:var(--blue);border-color:var(--blue);color:white}button:disabled{opacity:.35}button:focus-visible,a:focus-visible{outline:3px solid var(--blue);outline-offset:3px}.viewer{padding:20px;min-height:calc(100vh - 128px);display:grid;place-items:center}.deck{width:min(1400px,100%)}
.slide{display:none;position:relative;aspect-ratio:16/9;container-type:inline-size;background:white;border:1px solid var(--border);box-shadow:0 16px 48px #001f6110;overflow:hidden}.slide.active{display:block}.inner{padding:4.5cqw 5cqw}.eyebrow{font-size:1cqw;letter-spacing:.15em;font-weight:700;color:var(--blue);margin-bottom:1.6cqw}h1{font-size:3.8cqw;line-height:1.08;letter-spacing:-.035em;color:var(--navy);margin:0;max-width:85%}h1 br{display:block}.intro{font-size:1.6cqw;line-height:1.5;color:var(--body);max-width:77%;margin:1.6cqw 0 0}.cards{display:grid;gap:2cqw;margin-top:3.2cqw}.three{grid-template-columns:repeat(3,1fr)}.two{grid-template-columns:1fr 1fr}.card{padding:2.2cqw;border:1px solid var(--border);border-radius:1cqw;background:#f9fafb}.card h3{font-size:1.7cqw;color:var(--navy);margin:1.3cqw 0 0}.card p{font-size:1.5cqw;line-height:1.5;margin:1.2cqw 0 0;color:var(--body)}.card p:empty{display:none}.icon{display:block;width:3.8cqw;height:3.8cqw;stroke:var(--blue);stroke-width:2.5;fill:none;stroke-linecap:round;stroke-linejoin:round}.icon.small{width:3.1cqw;height:3.1cqw}.metrics{display:grid;grid-template-columns:repeat(3,1fr);gap:1.7cqw;margin-top:4cqw}.metrics article{padding:2.1cqw 1.8cqw;border-top:3px solid var(--blue);background:var(--soft);border-radius:.7cqw}.metrics b{display:block;font-size:4.5cqw;line-height:1.1;font-weight:700;letter-spacing:-.05em;color:var(--navy);white-space:nowrap}.metrics span{display:block;font-size:1.3cqw;color:var(--body);margin-top:1.1cqw}.producer-row{display:flex;gap:5cqw;margin-top:2.4cqw}.producer-row div{display:flex;gap:1.4cqw;align-items:center}.producer-row strong{font-size:1.45cqw;color:var(--navy)}.producer-row span{font-size:1.35cqw;color:var(--body)}.caption{font-size:1.5cqw;color:var(--body);line-height:1.5;margin:2.5cqw 0 0}.disclosure{font-size:1.25cqw;line-height:1.5;color:var(--navy);padding:1cqw 1.3cqw;background:var(--soft);border-left:3px solid var(--blue);margin:2cqw 0 0}.architecture{margin-top:3cqw}.lane{display:flex;align-items:center;gap:1.4cqw}.step{display:flex;flex-direction:column;gap:1.2cqw;justify-content:center;align-items:center;padding:1.9cqw 1cqw;border:1px solid var(--border);border-radius:.9cqw;background:var(--soft);min-width:0;flex:1}.step strong{font-size:1.35cqw;color:var(--navy);text-align:center;white-space:nowrap}.arrow{font-size:2.2cqw;color:var(--blue)}.branch{display:flex;gap:1.5cqw;align-items:center;justify-content:center;margin:1.4cqw auto 0;width:44%}.branch .step{padding:1.2cqw}.branch>span{font-size:2cqw;color:var(--blue)}.flow{display:flex;align-items:center;gap:.6cqw;margin-top:4.5cqw}.flow .step{padding:1.8cqw .3cqw}.flow .step strong{font-size:1.12cqw}.demo{font-size:2cqw;letter-spacing:-.015em;color:var(--navy);font-weight:700;margin-top:3cqw}.demo span{color:var(--blue);padding:0 .6cqw}.slide-footer{position:absolute;bottom:2.2cqw;left:5cqw;right:5cqw;display:flex;justify-content:space-between;align-items:center;border-top:1px solid var(--border);padding-top:1cqw;font-size:.85cqw;color:#29343d}.slide:before{content:'';position:absolute;inset:0 0 auto;height:.45cqw;background:var(--blue)}.status{font-size:13px;border-top:1px solid var(--border);color:var(--navy)}
@media(max-width:700px){.toolbar{padding:10px;flex-wrap:wrap}.controls{gap:5px}button,a{font-size:12px;padding:9px}.viewer{padding:6px;min-height:calc(100vh - 150px)}.status{padding:12px}.hint{display:none}}
@page{size:13.333in 7.5in;margin:0}@media print{body{background:white}.toolbar,.status{display:none}.viewer{display:block;padding:0;min-height:0}.deck{width:13.333in}.slide,.slide.active{display:block!important;width:13.333in;height:7.5in;aspect-ratio:auto;border:0;box-shadow:none;break-after:page;break-inside:avoid}.slide:last-child{break-after:auto}}
"""


def write_html():
    sections=[]
    for i,(label,title,intro,visual) in enumerate(DATA,1):
        sections.append(f'<section class="slide{" active" if i==1 else ""}" aria-label="Slide {i}"><div class="inner"><div class="eyebrow">{label}</div><h1>{title}</h1>'+ (f'<p class="intro">{intro}</p>' if intro else '')+visual+f'</div><footer class="slide-footer"><span>MERIDIAN · Fictional brokerage · Synthetic data</span><span>{i:02} / 06</span></footer></section>')
    page='<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Meridian | Six-slide solution</title><style>'+CSS+'</style></head><body><header class="toolbar"><span class="brand">MERIDIAN / THE SOLUTION</span><nav class="controls" aria-label="Presentation"><button id="prev">← Previous</button><button id="next" class="primary">Next →</button><button id="print">Print</button><a href="meridian-solution.pdf" download>PDF</a></nav></header><main class="viewer"><div class="deck">'+''.join(sections)+'</div></main><footer class="status"><span id="counter" aria-live="polite"></span><span class="hint">Arrow keys to navigate · Home / End · Print shows all slides</span></footer><script>let n=0;const s=[...document.querySelectorAll(".slide")];function show(i){n=Math.max(0,Math.min(5,i));s.forEach((e,j)=>e.classList.toggle("active",j===n));document.querySelector("#prev").disabled=n===0;document.querySelector("#next").disabled=n===5;document.querySelector("#counter").textContent=`Slide ${n+1} of 6`;history.replaceState(null,"",`#slide-${n+1}`)}document.querySelector("#prev").onclick=()=>show(n-1);document.querySelector("#next").onclick=()=>show(n+1);document.querySelector("#print").onclick=()=>print();document.addEventListener("keydown",e=>{if(e.altKey||e.ctrlKey||e.metaKey)return;if(["ArrowRight","ArrowLeft","Home","End"].includes(e.key)){e.preventDefault();show(e.key==="Home"?0:e.key==="End"?5:n+(e.key==="ArrowRight"?1:-1))}});show(Number(location.hash.split("-")[1]||1)-1);</script></body></html>'
    (HERE/'meridian-solution.html').write_text(page,encoding='utf-8')


def write_pptx():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    def text(slide,x,y,w,h,value,size=22,color=NAVY,bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        box.text_frame.word_wrap=True
        for j,value in enumerate(value.split('\n')):
            p=box.text_frame.paragraphs[0] if j==0 else box.text_frame.add_paragraph();p.text=value;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
    def panel(slide,x,y,w,h,title,body):
        box=slide.shapes.add_shape(Shape.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));box.fill.solid();box.fill.fore_color.rgb=RGBColor.from_string(SOFT);box.line.color.rgb=RGBColor.from_string(BORDER)
        text(slide,x+.2,y+.25,w-.4,.7,title,26,bold=True);text(slide,x+.2,y+1.1,w-.4,h-1.2,body,19,INK)
    content=[
        [('See the risk','Renewal priorities'),('Check the evidence','Sources and conflicts'),('Choose the next step','Producer judgement')],
        [('9','Authorized clients'),('2','Producers'),('$5.44M','Annual premium')],
        [('React + TypeScript','Browser workspace'),('FastAPI + Pydantic','Permissions + typed contracts'),('SQLite + models','Relational sources + reviewable AI')],
        [('Permissions → SQL FTS5','Only permitted evidence'),('Typed context → Models','Structured, grounded preparation'),('Citations → Brief','Canonical sources; labeled fallback')],
        [('Producer','Read brief; inspect sources; edit and manually copy draft.'),('Regional Leader','Prioritize renewals; open preparation; review database dashboard.')],
        [('2.37s','One observed Pinecrest response'),('1,460','Measured tokens'),('C05 + E05','Canonical citations')],
    ]
    captions={1:'Jordan Davis: 5 clients / $3.33M    ·    Priya Nair: 4 / $2.11M\nSynthetic finance: 12% commission; 35% paid loss; 100% retained demo outcomes.',3:'User and source hashes prevent stale reuse.',4:'No autonomous messages, policy changes, or coverage binding.',5:'One request, no latency guarantee. Synthetic data; two personas; no proven uplift.\nDemo: prioritize → inspect → prepare → copy.'}
    for i,(label,title,intro,_) in enumerate(DATA):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);text(slide,.6,.35,12,.35,label,12,BLUE,True);text(slide,.6,.95,12,1.2,title.replace('<br>','\n'),34,bold=True)
        if intro:text(slide,.6,2.15,12,.95,intro,21,INK)
        entries=content[i];width=(12.1-(len(entries)-1)*.25)/len(entries)
        for j,(heading,body) in enumerate(entries):panel(slide,.6+j*(width+.25),3.3,width,2.25,heading,body)
        if i in captions:text(slide,.6,5.85,12,.85,captions[i],16,INK)
        text(slide,.6,7.05,11,.25,'MERIDIAN · Fictional brokerage · Synthetic data',10,INK);text(slide,12,7.05,.7,.25,f'{i+1:02}/06',10,BLUE)
    prs.save(HERE/'meridian-solution.pptx')


def print_pdf():
    with sync_playwright() as p:
        browser=p.chromium.launch(channel='msedge',headless=True)
        page=browser.new_page(viewport={'width':1440,'height':900});page.goto((HERE/'meridian-solution.html').as_uri())
        page.locator('#next').click();assert page.locator('#counter').inner_text()=='Slide 2 of 6';page.keyboard.press('Home')
        page.screenshot(path=str(HERE/'html-preview.png'),full_page=True)
        page.pdf(path=str(HERE/'meridian-solution.pdf'),prefer_css_page_size=True,print_background=True)
        browser.close()
    pages=len(PdfReader(HERE/'meridian-solution.pdf').pages);assert pages==6,pages
    return pages


if __name__=='__main__':
    write_html();write_pptx();pages=print_pdf()
    result={'slides':6,'pdf_pages':pages,'html_navigation':'passed','provider_calls':0,'database_calls':0}
    (HERE/'artifact-check.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result))
